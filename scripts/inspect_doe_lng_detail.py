#!/usr/bin/env python3
"""Inspect official DOE LNG cargo-detail workbook; emit sheet/schema audit before aggregating."""
import io,json,sys,urllib.request
from pathlib import Path
from openpyxl import load_workbook
URL='https://www.energy.gov/sites/default/files/2026-10/3.%20U.S.%20LNG%20Exports%20and%20Re-Exports%20Details%20%28Jan%202016%20-%20Aug%202026%29.xlsx'
def main(path):
 req=urllib.request.Request(URL,headers={'User-Agent':'Mozilla/5.0 EPNewsResearch/1.0'})
 with urllib.request.urlopen(req,timeout=120) as r:raw=r.read()
 book=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
 out={'source':URL,'bytes':len(raw),'sheets':[],'warning':'Schema inventory only: no country or period totals are asserted until columns and units are independently validated'}
 for sh in book:
  preview=[]
  for i,row in enumerate(sh.iter_rows(values_only=True)):
   if i>=12:break
   preview.append([str(v)[:160] if v is not None else None for v in row[:25]])
  out['sheets'].append({'name':sh.title,'rows':sh.max_row,'columns':sh.max_column,'first_12_rows':preview})
 Path(path).write_text(json.dumps(out,indent=2)+'\n')
 print('DOE workbook',len(raw),'bytes;',len(out['sheets']),'sheets')
 for sh in out['sheets']:print(sh['name'],sh['rows'],'rows',sh['columns'],'cols')
if __name__=='__main__':main(sys.argv[1] if len(sys.argv)>1 else 'doe-lng-detail-schema.json')
