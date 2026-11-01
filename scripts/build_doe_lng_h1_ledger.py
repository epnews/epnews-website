#!/usr/bin/env python3
"""DOE official LNG exports: destination ledger for 2025/2026 H1, no re-export double count."""
import io,json,sys,urllib.request,datetime,collections
from pathlib import Path
from openpyxl import load_workbook
URL='https://www.energy.gov/sites/default/files/2026-10/3.%20U.S.%20LNG%20Exports%20and%20Re-Exports%20Details%20%28Jan%202016%20-%20Aug%202026%29.xlsx'
EU={'Austria','Belgium','Bulgaria','Croatia','Cyprus','Czech Republic','Denmark','Estonia','Finland','France','Germany','Greece','Hungary','Ireland','Italy','Latvia','Lithuania','Luxembourg','Malta','Netherlands','Poland','Portugal','Romania','Slovakia','Slovenia','Spain','Sweden'}
FACTOR=0.000028316846592
HEADERS=['Arrival/Departure Date','Activity','Gas Type','Mode of Transport','Country','Volume (MMCF)']
def ledger(rows):
 it=iter(rows);headers=list(next(it));missing=set(HEADERS)-set(headers)
 if missing:raise ValueError('Missing columns: '+str(missing))
 ix={h:headers.index(h) for h in HEADERS}
 out=collections.defaultdict(lambda:collections.defaultdict(lambda:collections.defaultdict(float)))
 counts=collections.Counter();errors=[]
 for n,row in enumerate(it,2):
  def cell(h):return row[ix[h]] if ix[h]<len(row) else None
  d=cell('Arrival/Departure Date')
  if isinstance(d,str):
   try:d=datetime.datetime.fromisoformat(d)
   except ValueError:continue
  if not isinstance(d,(datetime.date,datetime.datetime)) or d.year not in (2025,2026) or d.month>6:continue
  country=str(cell('Country') or '').strip()
  activity=str(cell('Activity') or '').strip().lower().replace('-','')
  mode=str(cell('Mode of Transport') or '').strip().lower()
  gas=str(cell('Gas Type') or '').strip().upper()
  if gas!='LNG' or mode!='vessel':continue
  if activity not in ('exports','reexports'):
   errors.append(f'row {n} unrecognized activity {activity!r}');continue
  volume=cell('Volume (MMCF)')
  try:volume=float(volume)
  except (ValueError,TypeError):errors.append(f'row {n} invalid volume');continue
  if volume<0:errors.append(f'row {n} negative volume');continue
  if not country:errors.append(f'row {n} missing destination');continue
  y=str(d.year)
  out[y][country][activity]+=volume
  counts[(y,activity)]+=1
 if errors:raise ValueError('Invalid DOE transaction rows: '+ '; '.join(errors[:12]))
 return out,counts
def main(path):
 req=urllib.request.Request(URL,headers={'User-Agent':'Mozilla/5.0 EPNewsResearch/1.0'})
 with urllib.request.urlopen(req,timeout=120) as r:raw=r.read()
 book=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
 sh=book['By Vessel and ISO Container']
 out,counts=ledger(sh.iter_rows(values_only=True))
 result={'source':URL,'periods':{},'unit':'bcm','conversion_bcm_per_mmcf':FACTOR,'scope':'DOE departure-date LNG vessel shipments, exports separately from re-exports','caveats':['DOE destination is not proof of arrival in the same month','Country may be declared destination rather than ultimate consumption','DOE departures and Eurostat imports differ in scope, reporting, and geography','No net economic benefit or influence inference']}
 for y in ('2025','2026'):
  destinations={}
  for country,activities in sorted(out[y].items()):
   destinations[country]={k+'_bcm':round(activities.get(k,0)*FACTOR,9) for k in ('exports','reexports')}
  eu_exports=sum(out[y][c].get('exports',0) for c in EU)*FACTOR
  eu_reexports=sum(out[y][c].get('reexports',0) for c in EU)*FACTOR
  result['periods'][y]={'destinations':destinations,'eu27_exports_bcm':round(eu_exports,9),'eu27_reexports_bcm':round(eu_reexports,9),'eu27_destination_countries_with_exports':sorted(c for c in EU if out[y][c].get('exports',0)>0),'export_transactions':counts[(y,'exports')],'reexport_transactions':counts[(y,'reexports')]}
 Path(path).write_text(json.dumps(result,indent=2)+'\n')
 for y,p in result['periods'].items():print(y,'EU27 exports bcm',p['eu27_exports_bcm'],'reexports bcm',p['eu27_reexports_bcm'],'export transactions',p['export_transactions'])
if __name__=='__main__':main(sys.argv[1] if len(sys.argv)>1 else 'doe-lng-h1-destination-ledger.json')
